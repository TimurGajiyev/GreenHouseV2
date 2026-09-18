"""CHP_BESS_dispatch_sim.xlsx: rebuild its arithmetic, then find what drives it.

The workbook is a forward rule simulation of one 24-hour day at 1-minute
resolution, annualised over 330 operating days. Its headline is that the battery
LOSES 97.7 million tenge a year and never pays back, which is the opposite of
CHP_BESS_model_v2.xlsx (6.3 years). This script:

  1. rebuilds every total on its Results sheet from its own inputs,
  2. isolates which assumption produces the sign flip,
  3. measures how much of its load variation lives below the hourly average,
     which is the part an hourly model -- ours and model_v2 alike -- cannot see.
"""

from __future__ import annotations

import io
import json
import os
import sys

import numpy as np
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
OUT = os.path.join(ROOT, "chp_bess_analysis")
XL = os.path.join(ROOT, "CHP_BESS_dispatch_sim.xlsx")


def load():
    wb = openpyxl.load_workbook(XL, data_only=True)
    inp = {}
    for r in wb["Inputs"].iter_rows(values_only=True):
        if r[0] and r[1] is not None and isinstance(r[1], (int, float)):
            inp.setdefault(str(r[0]).strip(), float(r[1]))
    res = {}
    for r in wb["Results"].iter_rows(values_only=True):
        if r[0] and isinstance(r[0], str):
            res[r[0].strip()] = (r[1], r[2], r[3])
    sim = wb["Sim"]
    hdr = [c.value for c in sim[1]]
    cols = {h: i for i, h in enumerate(hdr) if h}
    rows = list(sim.iter_rows(min_row=2, values_only=True))
    return inp, res, cols, rows


def main():
    inp, res, cols, rows = load()
    g = lambda k: res[k]
    n = lambda x: 0.0 if x is None else float(x)

    gas_price = inp["Gas cost per MWh of gas energy"]
    grid_mwh = inp["Grid import tariff"] if inp["Grid import tariff"] > 1000 else inp["Grid import tariff"] * 1000
    om_h = inp["O&M cost per engine running hour"]
    wear_h = inp["Wear cost of one start"]
    capex = inp["BESS installed capex"]
    days = inp["Operating days per year"]

    print("Inputs that decide the answer")
    print(f"  gas            {gas_price:>12,.2f} ₸/МВт·ч(газ)   ({inp['Gas price']:,.0f} ₸/1000 м³, LHV {inp['Gas lower heating value']} кВт·ч/м³)")
    print(f"  grid           {grid_mwh:>12,.0f} ₸/МВт·ч")
    print(f"  O&M            {om_h:>12,.0f} ₸ за МОТОЧАС агрегата   <-- не за кВт·ч")
    print(f"  start wear     {wear_h:>12,.0f} эквивалентных моточасов = {wear_h*om_h:>10,.0f} ₸ за пуск")
    print(f"  CAPEX          {capex:>12,.0f} ₸        дней в году {days:,.0f}")
    print()

    # ---- 1. rebuild every total -------------------------------------------
    print("1. Пересчёт их итогов из их же входов")
    checks = []
    for label, key, col, calc in (
        ("Gas cost, baseline", "Gas cost (KZT)", 0, n(g("Gas consumed (MWh, LHV)")[0]) * gas_price),
        ("Gas cost, BESS", "Gas cost (KZT)", 1, n(g("Gas consumed (MWh, LHV)")[1]) * gas_price),
        ("Grid cost, baseline", "Grid import cost (KZT)", 0, n(g("Grid import (MWh)")[0]) * grid_mwh),
        ("O&M baseline", "Engine O&M (KZT)", 0, n(g("Effective hours (running + start wear)")[0]) * om_h),
        ("O&M with BESS", "Engine O&M (KZT)", 1, n(g("Effective hours (running + start wear)")[1]) * om_h),
        ("Start wear hours", "Start wear (equivalent run hours)", 1, n(g("Engine starts")[1]) * wear_h),
    ):
        got, want = calc, n(g(key)[col])
        checks.append(abs(got - want) < 1.0)
        print(f"   {'OK ' if checks[-1] else 'XX '} {label:<24} {got:>14,.1f}  vs  {want:>14,.1f}")
    base = n(g("Total operating cost (KZT)")[0])
    bess = n(g("Total operating cost (KZT)")[1])
    day_saving = base - bess
    print(f"   {'OK ' if abs(day_saving*days - n(g('Saving per year (KZT)')[0])) < 2 else 'XX '} "
          f"{'annualised':<24} {day_saving*days:>14,.0f}  vs  {n(g('Saving per year (KZT)')[0]):>14,.0f}")
    print(f"   -> экономия за сутки {day_saving:,.0f} ₸, за год {day_saving*days:,.0f} ₸"
          f"  =>  {'нет окупаемости' if day_saving <= 0 else f'{capex/(day_saving*days):.1f} лет'}\n")

    # ---- 2. where the sign comes from -------------------------------------
    fuel_saving = n(g("Gas cost (KZT)")[0]) - n(g("Gas cost (KZT)")[1])
    grid_saving = n(g("Grid import cost (KZT)")[0]) - n(g("Grid import cost (KZT)")[1])
    hours_saved = n(g("Total fleet running hours")[0]) - n(g("Total fleet running hours")[1])
    starts = n(g("Engine starts")[1])
    print("2. Из чего складывается знак (за сутки)")
    print(f"   экономия газа                    {fuel_saving:>+12,.0f} ₸")
    print(f"   экономия на закупе из сети       {grid_saving:>+12,.0f} ₸")
    print(f"   экономия моточасов {hours_saved:>5.1f} ч       {hours_saved*om_h:>+12,.0f} ₸")
    print(f"   износ от {starts:.0f} пусков x {wear_h:.0f} ч         {-starts*wear_h*om_h:>+12,.0f} ₸   <-- перевешивает всё")
    print(f"   {'итого':<32} {fuel_saving+grid_saving+hours_saved*om_h-starts*wear_h*om_h:>+12,.0f} ₸\n")

    # ---- 3. the break-even on the decisive placeholder ---------------------
    energy_saving = fuel_saving + grid_saving
    print("3. Порог по единственному решающему допущению")
    be_wear = (energy_saving + hours_saved * om_h) / (starts * om_h) if starts else float("inf")
    print(f"   износ пуска, при котором проект выходит в ноль: {be_wear:.2f} эквивалентных моточасов")
    print(f"   их допущение: {wear_h:.0f} ч (диапазон OEM в самой книге: 10-20 ч)")
    print(f"   модель v2 оценивает пуск в 15 000 ₸ = {15000/om_h:.2f} моточаса -- в {wear_h*om_h/15000:.1f} раза дешевле\n")
    print("   окупаемость при разных допущениях:")
    for lbl, sav in (
        (f"как в книге (износ {wear_h:.0f} ч, O&M за моточас)", energy_saving + hours_saved*om_h - starts*wear_h*om_h),
        ("контракт O&M за МВт·ч (ставка за моточас = 0)", energy_saving),
        (f"износ пуска {15000/om_h:.2f} ч, как в model_v2", energy_saving + hours_saved*om_h - starts*(15000/om_h)*om_h),
        ("износ пуска 0", energy_saving + hours_saved*om_h),
    ):
        yr = sav * days
        print(f"      {lbl:<48} {yr:>+14,.0f} ₸/год  ->  "
              f"{'нет окупаемости' if yr <= 0 else f'{capex/yr:.1f} лет'}")
    print()

    # ---- 4. what an hourly model cannot see -------------------------------
    ci = cols.get("Site load (MW)")
    load_min = np.array([r[ci] for r in rows if r[ci] is not None], dtype=float)
    hourly = load_min[:1440].reshape(24, 60).mean(axis=1)
    resid = load_min[:1440].reshape(24, 60) - hourly[:, None]
    print("4. Что теряет почасовая модель (наша и model_v2 одинаково)")
    print(f"   нагрузка: {len(load_min)} минут, среднее {load_min.mean():.3f} МВт, "
          f"мин {load_min.min():.3f}, макс {load_min.max():.3f}")
    print(f"   почасовое среднее: мин {hourly.min():.3f}, макс {hourly.max():.3f} МВт")
    print(f"   внутричасовое отклонение: sd {resid.std():.3f} МВт, "
          f"размах {resid.min():.3f}..{resid.max():.3f} МВт")
    print(f"   пик минутный {load_min.max():.3f} против пика часового {hourly.max():.3f} МВт "
          f"-> почасовая модель занижает пик на {(load_min.max()-hourly.max())/load_min.max():.1%}")
    os.makedirs(OUT, exist_ok=True)
    io.open(os.path.join(OUT, "dispatch_sim_day_hourly.json"), "w", encoding="utf-8").write(
        json.dumps({"hourly_mw": hourly.tolist(), "minute_mw": load_min.tolist()}))
    print(f"\n   почасовой профиль их суток сохранён -> chp_bess_analysis/dispatch_sim_day_hourly.json")


if __name__ == "__main__":
    main()
